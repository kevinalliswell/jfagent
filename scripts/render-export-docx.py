#!/usr/bin/env python3
"""Render ExportPayloadV1 JSON into a deterministic DOCX requirement sheet."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


CORPORATE_BLUE = "1F4E79"
RISK_RED = "C00000"
RISK_FILL = "FCE4D6"
PRICE_FILL = "FFF2CC"
PRICE_TEXT = "9C6500"
LIGHT_FILL = "F2F4F7"
TABLE_WIDTH_CM = 16.0


def text(value: Any, fallback: str = "待确认") -> str:
    if value is None or value == "":
        return fallback
    return str(value)


def number_text(value: Any, unit: str = "") -> str:
    if value is None or value == "":
        return "待确认"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"{value}{unit}"


def sanitize_filename(value: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|\\s]+", "_", value).strip("_")
    return cleaned or "jfagent_export"


def set_east_asian_font(run, font_name: str = "Microsoft YaHei") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_style_font(style, font_name: str, size_pt: float, color: str | None = None, bold: bool | None = None) -> None:
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.font.size = Pt(size_pt)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def style_paragraph(style, before: int = 0, after: int = 6, line_spacing: float = 1.25) -> None:
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def configure_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    styles = doc.styles
    set_style_font(styles["Normal"], "Microsoft YaHei", 10.5)
    style_paragraph(styles["Normal"], after=6, line_spacing=1.25)

    for name, size, color, bold, before, after in [
        ("Heading 1", 16, CORPORATE_BLUE, True, 16, 8),
        ("Heading 2", 13, "000000", True, 12, 6),
        ("Heading 3", 11, "000000", True, 8, 4),
    ]:
        set_style_font(styles[name], "Microsoft YaHei", size, color, bold)
        style_paragraph(styles[name], before=before, after=after, line_spacing=1.15)

    for name, size, color, bold, before, after in [
        ("CoverCustomer", 16, "666666", False, 0, 8),
        ("CoverTitle", 24, CORPORATE_BLUE, True, 0, 8),
        ("CoverSubtitle", 14, "333333", False, 0, 18),
        ("BodyText", 10.5, "000000", False, 0, 6),
        ("MutedText", 9, "666666", False, 0, 4),
    ]:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(style, "Microsoft YaHei", size, color, bold)
        style_paragraph(style, before=before, after=after, line_spacing=1.25)

    price_style = styles.add_style("PricePlaceholder", WD_STYLE_TYPE.CHARACTER)
    set_style_font(price_style, "Arial", 7.5, PRICE_TEXT, False)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("JF Agent Workstation - 内部售前需求表")
    set_east_asian_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return doc


def table_pr(table, total_width_cm: float = TABLE_WIDTH_CM) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(total_width_cm * 567)))

    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "0")

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")]:
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, cm: float) -> None:
    cell.width = Cm(cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(cm * 567)))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "BFBFBF", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ["top", "left", "bottom", "right"]:
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_text(cell, value: Any, *, bold: bool = False, color: str | None = None, size: float = 9.5, align=WD_ALIGN_PARAGRAPH.LEFT, price: bool = False) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.alignment = align
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.1
    para.clear()
    run = para.add_run(text(value, ""))
    set_east_asian_font(run)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if price:
        run.style = "PricePlaceholder"
        set_cell_shading(cell, PRICE_FILL)
    set_cell_borders(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float], *, header_fill: str = CORPORATE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table_pr(table, sum(widths))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, header_fill)
        set_cell_text(cell, header, bold=True, color="FFFFFF" if header_fill == CORPORATE_BLUE else "000000", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            set_cell_width(cell, widths[index])
            set_cell_text(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if widths[index] <= 1.4 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph("")


def add_price_text(cell, price_payload: dict[str, Any]) -> None:
    placeholder = price_payload.get("placeholder_text") or "[Please manually enter your local channel price here]"
    set_cell_text(cell, placeholder, size=7.0, price=True)


def add_bom_table(doc: Document, lines: list[dict[str, Any]], title: str) -> None:
    if not lines:
        return
    doc.add_heading(title, level=2)
    headers = ["设备名称", "规格建议", "数量", "单位", "单价", "合价", "备注"]
    widths = [2.6, 3.4, 1.1, 1.0, 2.4, 2.4, 3.1]
    table = doc.add_table(rows=1, cols=len(headers))
    table_pr(table, sum(widths))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, CORPORATE_BLUE)
        set_cell_text(cell, header, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in lines:
        cells = table.add_row().cells
        values = [
            line.get("item_name"),
            line.get("specification"),
            line.get("quantity"),
            line.get("unit"),
            line.get("unit_price", {}),
            line.get("subtotal_price", {}),
            line.get("remark"),
        ]
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_width(cell, widths[index])
            if index in (4, 5):
                add_price_text(cell, value if isinstance(value, dict) else {})
            else:
                set_cell_text(cell, value, size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER if index in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph("")


def add_key_value_table(doc: Document, rows: list[tuple[str, Any]], widths: list[float] | None = None) -> None:
    add_table(doc, ["字段", "内容"], [[label, text(value)] for label, value in rows], widths or [4.2, 11.8])


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run(item)
        set_east_asian_font(run)
        run.font.size = Pt(10.5)


def add_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    para = doc.add_paragraph()
    run = para.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "请在 Word 中更新目录域以显示页码。"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    doc.add_page_break()


def render_cover(doc: Document, payload: dict[str, Any]) -> None:
    doc.core_properties.title = payload.get("project_name", "机房需求表")
    doc.core_properties.subject = "Data Center Requirement Sheet"
    doc.core_properties.author = payload.get("export_metadata", {}).get("generated_by", "Data Center Pre-sales AI Agent")

    for _ in range(5):
        doc.add_paragraph("")
    customer = doc.add_paragraph(style="CoverCustomer")
    customer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    customer.add_run(text(payload.get("customer_name"), "客户名称待确认"))
    title = doc.add_paragraph(style="CoverTitle")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(text(payload.get("project_name"), "数据中心机房建设项目"))
    subtitle = doc.add_paragraph(style="CoverSubtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("数据中心机房建设需求表")

    metadata = payload.get("export_metadata", {})
    add_key_value_table(
        doc,
        [
            ("文档版本", payload.get("version", "v1")),
            ("生成时间", metadata.get("generated_at")),
            ("方案类型", "售前需求表 / Requirement Sheet"),
            ("模板版本", metadata.get("template_version")),
            ("价格口径", metadata.get("pricing_mode")),
            ("编制单位", "Pre-sales AI Copilot"),
        ],
        [4.2, 10.3],
    )
    if metadata.get("calculation_status") == "blocked":
        warning = doc.add_paragraph(style="BodyText")
        warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = warning.add_run("注意：当前测算状态为 blocked，本文件仅可作为草稿，不能作为施工安全结论。")
        set_east_asian_font(run)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(RISK_RED)
    doc.add_page_break()


INDUSTRY_PITCH = {
    "medical": "医疗行业机房承载 HIS、PACS、LIS 等核心业务系统，建设方案应优先保障连续运行、故障可维护性和设施安全边界。本需求表以已确认的面积、机柜、UPS 后备时间和现场楼层信息为基础，形成供深化设计、供应商询价和内部评审使用的统一口径。",
    "education": "教育行业机房通常兼顾教学、科研、办公与校园网络支撑，建设方案需要在可靠性、可扩展性和预算可控之间取得平衡。本需求表先固化关键规模参数，便于后续形成清晰的设备与施工范围。",
    "government": "政务机房建设应关注公共服务连续性、标准化验收和运维安全。本需求表以售前阶段可确认的信息为边界，先形成可评审、可询价、可继续深化的基础技术口径。",
    "unknown": "数据中心机房建设应围绕供配电、制冷、结构承重、综合布线、消防与运维可维护性形成统一技术口径。本需求表用于把早期零散线索整理为后续深化设计和报价工作的基础输入。",
}


def render_chapter_1(doc: Document, payload: dict[str, Any]) -> None:
    doc.add_heading("第一章 项目概况与行业建设背景", level=1)
    doc.add_heading("1.1 项目基本信息", level=2)
    industry = payload.get("customer_industry", "unknown")
    project_type = payload.get("project_type", "unknown")
    add_key_value_table(
        doc,
        [
            ("客户名称", payload.get("customer_name")),
            ("项目名称", payload.get("project_name")),
            ("行业类型", industry),
            ("项目类型", project_type),
            ("机房面积", number_text(payload.get("room_area_m2"), " m2")),
            ("所在楼层", number_text(payload.get("room_floor"), " 楼")),
            ("机柜数量", number_text(payload.get("rack_count"), " 台")),
            ("UPS后备时间", number_text(payload.get("ups_backup_time_minutes"), " 分钟")),
            ("冗余模式", payload.get("redundancy_mode")),
        ],
    )
    doc.add_heading("1.2 行业背景与建设必要性", level=2)
    para = doc.add_paragraph(style="BodyText")
    para.add_run(INDUSTRY_PITCH.get(industry, INDUSTRY_PITCH["unknown"]))
    doc.add_heading("1.3 建设目标", level=2)
    add_bullets(
        doc,
        [
            "提升机房供配电系统可靠性与可维护性。",
            "保障核心 IT 设备连续运行与散热安全。",
            "形成规范化、可验收、可扩展的数据中心基础设施环境。",
            "为后续报价、深化设计、施工组织和验收交付提供统一技术口径。",
        ],
    )


def render_chapter_2(doc: Document, payload: dict[str, Any]) -> None:
    doc.add_heading("第二章 土建条件与 Steel Structure Load加固", level=1)
    table = doc.add_table(rows=1, cols=1)
    table_pr(table, TABLE_WIDTH_CM)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, RISK_FILL)
    set_cell_borders(cell, RISK_RED, "12")
    set_cell_text(
        cell,
        "Structural Loading Deficit Risk: 机房位于二层及以上，且 UPS 后备时间达到或超过 120 分钟。电池柜及 UPS 设备重量可能导致楼板承重不足，必须进行结构荷载复核。",
        bold=True,
        color=RISK_RED,
        size=10.5,
    )
    doc.add_paragraph("")
    doc.add_heading("2.1 触发条件", level=2)
    add_key_value_table(
        doc,
        [
            ("所在楼层", number_text(payload.get("room_floor"), " 楼")),
            ("UPS后备时间", number_text(payload.get("ups_backup_time_minutes"), " 分钟")),
            ("触发规则", "RULE_FLOOR_LOADING"),
            ("风险等级", "P0_BLOCKER"),
        ],
    )
    doc.add_heading("2.2 结构复核要求", level=2)
    add_bullets(
        doc,
        [
            "复核原建筑结构设计荷载与实际设备集中荷载。",
            "计算 UPS 主机、电池柜、电池组及底座的静载和集中荷载。",
            "必要时设置钢结构承重平台、分布梁或楼板加固措施。",
            "结构复核及加固方案应由业主或具备资质的结构专业单位确认。",
        ],
    )


def render_chapter_3(doc: Document, payload: dict[str, Any]) -> None:
    calc = payload.get("calculation_outputs", {})
    doc.add_heading("第三章 供配电与UPS工程方案", level=1)
    doc.add_heading("3.1 UPS容量计算结果", level=2)
    add_table(
        doc,
        ["参数", "数值", "单位", "备注"],
        [
            ["机柜数量", payload.get("rack_count"), "台", "用于 IT 负载估算"],
            ["单柜平均功率", calc.get("avg_power_per_rack_kw"), "kW/柜", "未确认时使用默认值并标记为暂估"],
            ["总IT负载", calc.get("total_it_load_kw"), "kW", "Rack Count * Avg Power per Rack"],
            ["冗余模式", payload.get("redundancy_mode"), "-", "N / N+1 / 2N"],
            ["UPS冗余系数", calc.get("redundancy_factor"), "-", "由冗余模式映射"],
            ["UPS原始计算容量", calc.get("required_ups_capacity_kva_raw"), "kVA", "未向下取整"],
            ["推荐UPS容量", calc.get("recommended_ups_capacity_kva"), "kVA", "按标准容量向上取整"],
        ],
        [4.1, 2.3, 2.0, 7.6],
    )
    bom = payload.get("bom", {})
    add_bom_table(doc, bom.get("ups", []), "3.2 UPS配置BOM")
    add_bom_table(doc, bom.get("battery", []), "3.3 电池后备时间与配置记录")
    add_bom_table(doc, bom.get("pdu", []), "3.4 PDU与配电回路配置")
    doc.add_heading("3.5 施工与验收要点", level=2)
    add_bullets(
        doc,
        [
            "UPS 输入、输出及旁路回路应按设计容量和现场配电条件进行复核。",
            "电池开关、线缆截面、端子压接和标识应满足安全和维护要求。",
            "PDU 回路编号、负载分配和机柜位置应与最终机柜平面图一致。",
            "投运前应完成市电/电池切换、旁路切换、告警联动和带载测试。",
        ],
    )


def render_chapter_4(doc: Document, payload: dict[str, Any]) -> None:
    calc = payload.get("calculation_outputs", {})
    doc.add_heading("第四章 精密空调与新风系统方案", level=1)
    doc.add_heading("4.1 热负荷与制冷量计算", level=2)
    add_table(
        doc,
        ["参数", "数值", "单位", "备注"],
        [
            ["机房面积", payload.get("room_area_m2"), "m2", "用于环境热负荷估算"],
            ["房间热密度", calc.get("room_thermal_density_kw_per_m2"), "kW/m2", "默认值需标记为暂估"],
            ["房间热负荷", calc.get("room_thermal_load_kw"), "kW", "Room Area * Room Thermal Density"],
            ["IT设备热负荷", calc.get("total_it_load_kw"), "kW", "通常按 IT 负载等效热负荷"],
            ["安全系数", calc.get("cooling_safety_margin_factor"), "-", "默认 1.15"],
            ["计算制冷量", calc.get("required_cooling_capacity_kw_raw"), "kW", "未向下取整"],
            ["推荐机型", calc.get("recommended_precision_ac_model_kw"), "kW", "映射至厂家离散机型"],
        ],
        [4.1, 2.3, 2.0, 7.6],
    )
    add_bom_table(doc, payload.get("bom", {}).get("cooling", []), "4.2 精密空调配置建议")
    doc.add_heading("4.3 N+1冗余配置说明", level=2)
    if calc.get("cooling_redundancy") == "N+1":
        add_key_value_table(
            doc,
            [
                ("计算所需制冷量", number_text(calc.get("required_cooling_capacity_kw_raw"), " kW")),
                ("单台推荐机型", number_text(calc.get("recommended_precision_ac_model_kw"), " kW")),
                ("运行机数量N", "1 台，按首版售前默认口径"),
                ("备用机数量+1", "1 台"),
                ("合计配置数量", "2 台"),
            ],
        )
    else:
        doc.add_paragraph("当前制冷冗余口径待确认，正式方案需结合冷热通道组织、室外机位置和品牌选型进行深化。", style="BodyText")
    doc.add_heading("4.4 新风与气流组织要求", level=2)
    add_table(
        doc,
        ["检查项", "要求", "状态"],
        [
            ["新风引入", "按机房正压、人员维护及规范要求设置", "待现场确认"],
            ["防尘过滤", "新风入口应设置过滤措施", "待现场确认"],
            ["防火阀", "穿越防火分区时按规范设置", "待深化"],
            ["冷凝水排放", "排水坡度、防倒灌、防渗漏", "待现场确认"],
            ["气流组织", "避免冷热短路，优先保障机柜进风侧", "待深化"],
        ],
        [3.0, 9.4, 3.6],
    )


def render_chapter_5(doc: Document, payload: dict[str, Any]) -> None:
    doc.add_heading("第五章 动环、消防、布线与装修范围", level=1)
    scope = payload.get("scope", {})
    rows = [[key, "纳入范围" if value else "暂未纳入", "以深化设计和现场踏勘为准"] for key, value in scope.items()]
    add_table(doc, ["范围项", "状态", "备注"], rows, [5.0, 3.0, 8.0])


def render_chapter_6(doc: Document, payload: dict[str, Any]) -> None:
    doc.add_heading("第六章 风险清单与待确认事项", level=1)
    risks = payload.get("risk_flags", [])
    if risks:
        add_table(
            doc,
            ["风险ID", "等级", "说明", "触发字段"],
            [[risk.get("id"), risk.get("level"), risk.get("text"), ", ".join(risk.get("trigger_fields", []))] for risk in risks],
            [3.2, 2.1, 7.7, 3.0],
        )
    else:
        doc.add_paragraph("当前未触发强制风险，但仍需完成现场条件复核。", style="BodyText")
    open_items = payload.get("open_items", [])
    if open_items:
        doc.add_heading("6.1 待确认事项", level=2)
        add_table(
            doc,
            ["字段/事项", "当前值", "原因", "建议跟进"],
            [
                [
                    item.get("label"),
                    item.get("current_value"),
                    item.get("reason"),
                    item.get("suggested_followup"),
                ]
                for item in open_items
            ],
            [3.0, 3.0, 3.0, 7.0],
        )


def render_chapter_7(doc: Document, payload: dict[str, Any]) -> None:
    doc.add_heading("第七章 商务价格占位附录", level=1)
    commercial = payload.get("commercial", {})
    para = doc.add_paragraph(style="BodyText")
    para.add_run(text(commercial.get("disclaimer"), "价格需由授权渠道或项目负责人人工确认。"))
    all_lines: list[dict[str, Any]] = []
    for lines in payload.get("bom", {}).values():
        if isinstance(lines, list):
            all_lines.extend(lines)
    add_bom_table(doc, all_lines, "7.1 汇总设备与施工占位清单")


def render_document(payload: dict[str, Any], output_path: Path) -> dict[str, Any]:
    doc = configure_document()
    render_cover(doc, payload)
    add_toc(doc)

    chapter_ids = [chapter.get("id") for chapter in payload.get("chapter_plan", [])]
    render_chapter_1(doc, payload)
    if "CHAPTER_2_CIVIL_STRUCTURAL_REINFORCEMENT" in chapter_ids:
        render_chapter_2(doc, payload)
    if "CHAPTER_3_POWER_UPS" in chapter_ids:
        render_chapter_3(doc, payload)
    if "CHAPTER_4_PRECISION_AC_FRESH_AIR" in chapter_ids:
        render_chapter_4(doc, payload)
    if "CHAPTER_5_MONITORING_FIRE_CABLING_FITOUT" in chapter_ids:
        render_chapter_5(doc, payload)
    if "CHAPTER_6_RISK_REGISTER" in chapter_ids:
        render_chapter_6(doc, payload)
    render_chapter_7(doc, payload)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return audit_docx(output_path, payload)


def audit_docx(path: Path, payload: dict[str, Any]) -> dict[str, Any]:
    checks: list[dict[str, str]] = []
    status = "passed"
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
            names = set(archive.namelist())
        checks.append({"id": "docx_zip_opened", "status": "passed", "message": "DOCX package is readable."})
        for required in ["[Please manually enter your local channel price here]", PRICE_FILL, "TOC"]:
            if required in xml:
                checks.append({"id": f"contains_{required[:12]}", "status": "passed", "message": f"Found {required}."})
            else:
                checks.append({"id": f"contains_{required[:12]}", "status": "warning", "message": f"Could not find {required} in document XML."})
                status = "warning"
        if "word/styles.xml" in names:
            checks.append({"id": "styles_present", "status": "passed", "message": "Styles part is present."})
        if payload.get("render_flags", {}).get("chapter_2_required") and "Steel Structure Load" not in xml:
            checks.append({"id": "forced_chapter_present", "status": "blocked", "message": "Forced structural chapter missing."})
            status = "blocked"
        elif payload.get("render_flags", {}).get("chapter_2_required"):
            checks.append({"id": "forced_chapter_present", "status": "passed", "message": "Forced structural chapter present."})
    except Exception as exc:  # pragma: no cover - reported to CLI.
        return {
            "status": "blocked",
            "checks": [{"id": "docx_package_read", "status": "blocked", "message": str(exc)}],
        }
    return {"status": status, "checks": checks}


def main() -> int:
    parser = argparse.ArgumentParser(description="Render ExportPayloadV1 JSON into DOCX.")
    parser.add_argument("payload_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--audit-json", type=Path)
    args = parser.parse_args()

    payload = json.loads(args.payload_json.read_text(encoding="utf-8"))
    audit = render_document(payload, args.output_docx)
    if args.audit_json:
        args.audit_json.parent.mkdir(parents=True, exist_ok=True)
        args.audit_json.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "output_docx": str(args.output_docx), "audit": audit}, ensure_ascii=False))
    return 0 if audit.get("status") != "blocked" else 2


if __name__ == "__main__":
    sys.exit(main())
