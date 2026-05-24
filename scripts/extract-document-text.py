#!/usr/bin/env python3
import argparse
import subprocess
import sys
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree import ElementTree


WORD_TEXT_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
WORD_TAB_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tab"
WORD_BREAK_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"
WORD_PARAGRAPH_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"


def normalize_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\r", "\n").split("\n")]
    compact = []
    previous_blank = False
    for line in lines:
        if not line:
            if not previous_blank:
                compact.append("")
            previous_blank = True
            continue
        compact.append(line)
        previous_blank = False
    return "\n".join(compact).strip()


def extract_docx_with_zip(file_path: Path) -> str:
    parts = []
    with zipfile.ZipFile(file_path) as archive:
        xml_names = [
            name
            for name in archive.namelist()
            if name.startswith("word/")
            and name.endswith(".xml")
            and (
                name == "word/document.xml"
                or name.startswith("word/header")
                or name.startswith("word/footer")
            )
        ]
        for name in sorted(xml_names):
            root = ElementTree.fromstring(archive.read(name))
            paragraph_parts = []
            for node in root.iter():
                if node.tag == WORD_TEXT_TAG and node.text:
                    paragraph_parts.append(node.text)
                elif node.tag == WORD_TAB_TAG:
                    paragraph_parts.append(" ")
                elif node.tag == WORD_BREAK_TAG:
                    paragraph_parts.append("\n")
                elif node.tag == WORD_PARAGRAPH_TAG and paragraph_parts:
                    paragraph_parts.append("\n")
            parts.append("".join(paragraph_parts))
    return normalize_text("\n".join(parts))


def extract_docx(file_path: Path) -> str:
    try:
        import docx  # type: ignore

        document = docx.Document(str(file_path))
        parts = []
        parts.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = normalize_text("\n".join(parts))
        if text:
            return text
    except Exception:
        pass
    return extract_docx_with_zip(file_path)


def extract_pdf_with_pdfplumber(file_path: Path) -> str:
    import pdfplumber  # type: ignore

    parts = []
    with pdfplumber.open(str(file_path)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            if page_text.strip():
                parts.append(page_text)
    return normalize_text("\n\n".join(parts))


def extract_pdf_with_pypdf(file_path: Path) -> str:
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(str(file_path))
    parts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text)
    return normalize_text("\n\n".join(parts))


def extract_pdf_with_pdftotext(file_path: Path) -> str:
    with TemporaryDirectory() as tmp:
        output_path = Path(tmp) / "out.txt"
        subprocess.run(
            ["pdftotext", "-layout", str(file_path), str(output_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        return normalize_text(output_path.read_text(encoding="utf-8", errors="ignore"))


def extract_pdf(file_path: Path) -> str:
    errors = []
    for extractor in (extract_pdf_with_pdfplumber, extract_pdf_with_pypdf, extract_pdf_with_pdftotext):
        try:
            text = extractor(file_path)
            if text:
                return text
        except Exception as exc:
            errors.append(f"{extractor.__name__}: {exc}")
    raise RuntimeError("; ".join(errors))


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract text from DOCX and PDF files for local RAG indexing.")
    parser.add_argument("file", type=Path)
    args = parser.parse_args()

    file_path = args.file
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".docx":
            text = extract_docx(file_path)
        elif suffix == ".pdf":
            text = extract_pdf(file_path)
        else:
            raise ValueError(f"Unsupported extension: {suffix}")
    except Exception as exc:
        print(f"Failed to extract {file_path}: {exc}", file=sys.stderr)
        return 2

    sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
